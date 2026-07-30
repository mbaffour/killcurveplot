"""
Figure construction and export.

One Plotly figure builder drives BOTH the interactive on-screen plot and every
exported file, so what you see is what you get — the R app rendered screen and
export through different paths, which is how they drift apart.
"""
from __future__ import annotations

import io
import math

import numpy as np

__all__ = ["PALETTES", "SHAPES", "DASHES", "build_figure", "figure_to_bytes",
           "cumulative_frames", "export_pptx", "export_gif", "render_images"]

# ── Palettes ────────────────────────────────────────────────────────────────
# Okabe-Ito is the published colourblind-safe 8-colour set; the others match
# the R app's options so saved settings port across.
PALETTES = {
    "custom": [],
    "colorblind": ["#000000", "#E69F00", "#56B4E9", "#009E73",
                   "#F0E442", "#0072B2", "#D55E00", "#CC79A7"],
    "viridis": ["#440154", "#46327E", "#365C8D", "#277F8E",
                "#1FA187", "#4AC16D", "#A0DA39", "#FDE725"],
    "plasma": ["#0D0887", "#5302A3", "#8B0AA5", "#B83289",
               "#DB5C68", "#F48849", "#FEBC2A", "#F0F921"],
    "publication": ["#1B2A41", "#C1292E", "#235789", "#F1D302",
                    "#4E8098", "#7A9E7E", "#A44A3F", "#5B5F97"],
    "rainbow": ["#E6194B", "#F58231", "#FFE119", "#3CB44B",
                "#42D4F4", "#4363D8", "#911EB4", "#F032E6"],
    "gray": ["#000000", "#333333", "#555555", "#777777",
             "#999999", "#BBBBBB", "#CCCCCC", "#DDDDDD"],
    "warm": ["#7C1D06", "#A63603", "#E6550D", "#FD8D3C",
             "#FDBE85", "#B30000", "#D7301F", "#EF6548"],
    "cool": ["#08306B", "#08519C", "#2171B5", "#4292C6",
             "#6BAED6", "#9ECAE1", "#C6DBEF", "#DEEBF7"],
}

# Plotly marker symbols, labelled the way a ggplot user expects.
SHAPES = ["circle", "square", "diamond", "triangle-up", "triangle-down",
          "cross", "x", "star", "hexagon", "pentagon",
          "circle-open", "square-open", "diamond-open", "triangle-up-open"]

DASHES = ["solid", "dash", "dot", "dashdot", "longdash", "longdashdot"]


def palette_for(name, n):
    """n colours from a named palette, cycling if the palette is shorter."""
    base = PALETTES.get(name) or PALETTES["colorblind"]
    if not base:
        base = PALETTES["colorblind"]
    return [base[i % len(base)] for i in range(max(n, 0))]


_NORM_LABEL = {"t0": "normalised to t₀", "max": "normalised to max",
               "zscore": "z-score", "control": "relative to control"}


def transform_note(cfg):
    """Short human description of any active data transform, or ''.

    Used in two places on purpose: the Y-axis title (so an exported figure can
    never be mistaken for raw signal) and the UI badge.
    """
    bits = []
    n = cfg.get("normalize")
    if n and n not in ("none", ""):
        bits.append(_NORM_LABEL.get(n, str(n)))
    try:
        w = int(cfg.get("smooth_window", 1) or 1)
    except (TypeError, ValueError):
        w = 1
    if w > 1:
        bits.append(f"smoothed ({w}-pt)")
    return " · ".join(bits)


def _axis_kwargs(cfg, which):
    """Translate our flat settings dict into a Plotly axis spec."""
    p = f"{which}_"
    stype = cfg.get(p + "scale_type", "linear")
    label = cfg.get(p + "axis_label", "Time" if which == "x" else "OD")
    if which == "y":
        # A transformed axis must say so ON THE FIGURE — the plot outlives the
        # screen it was configured on (PNG in a slide deck, PDF in a thesis).
        note = transform_note(cfg)
        if note:
            label = f"{label} — {note}"
    ax = {
        "title": {"text": label,
                  "font": {"size": cfg.get("axis_label_font_size", 16),
                           "family": cfg.get("font_family", "Arial"),
                           "color": cfg.get("axis_color", "#000000")}},
        "showgrid": bool(cfg.get(f"show_major_gridlines", False)),
        "gridcolor": cfg.get("major_gridline_color", "#E6E6E6"),
        "gridwidth": cfg.get("major_gridline_size", 0.5),
        "showline": True,
        "linecolor": cfg.get("axis_color", "#000000"),
        "linewidth": cfg.get("axis_line_width", 1.2),
        "ticks": "outside",
        "tickcolor": cfg.get("axis_color", "#000000"),
        "tickfont": {"size": cfg.get("axis_text_font_size", 14),
                     "family": cfg.get("font_family", "Arial"),
                     "color": cfg.get("axis_color", "#000000")},
        "zeroline": False,
        "mirror": bool(cfg.get("panel_border", True)),
        "automargin": True,
    }
    if stype == "log":
        ax["type"] = "log"
    elif stype == "reverse":
        ax["autorange"] = "reversed"
    else:
        ax["type"] = "linear"

    if which == "x" and cfg.get("axis_text_angle"):
        ax["tickangle"] = cfg.get("axis_text_angle", 0)

    if cfg.get(p + "minor_ticks"):
        ax["minor"] = {"ticks": "outside", "ticklen": 3,
                       "showgrid": bool(cfg.get("show_minor_gridlines", False)),
                       "gridcolor": cfg.get("minor_gridline_color", "#F2F2F2")}

    if cfg.get(p + "tick_interval"):
        try:
            ax["dtick"] = float(cfg[p + "tick_interval"])
        except (TypeError, ValueError):
            pass

    lo, hi = cfg.get(p + "min"), cfg.get(p + "max")
    if lo not in (None, "") and hi not in (None, ""):
        try:
            lo, hi = float(lo), float(hi)
            # Plotly log axes take EXPONENTS, not data values. Passing raw data
            # values here silently produces a wildly wrong range.
            ax["range"] = [math.log10(max(lo, 1e-12)), math.log10(max(hi, 1e-12))] \
                if stype == "log" else [lo, hi]
        except (TypeError, ValueError):
            pass
    return ax


def build_figure(agg, cfg, styles=None, analysis_overlays=None, samples=None):
    """Build the Plotly figure.

    agg      : DataFrame from data.aggregate (sample, time, mean, err, n)
    cfg      : flat settings dict (see DEFAULTS in server.py)
    styles   : {sample: {color, shape, dash, label, visible}}
    """
    import plotly.graph_objects as go

    cfg = cfg or {}
    styles = styles or {}
    fig = go.Figure()

    names = samples if samples is not None else (
        list(dict.fromkeys(agg["sample"].tolist())) if agg is not None and len(agg) else [])
    pal = palette_for(cfg.get("color_palette", "colorblind"), len(names))

    err_mode = cfg.get("error_display_mode", "bar")     # bar | ribbon | none
    show_pts = bool(cfg.get("show_points", True))
    show_line = bool(cfg.get("show_lines", True))
    mult = float(cfg.get("error_multiplier", 1) or 1)

    for i, name in enumerate(names):
        st = styles.get(name, {})
        if st.get("visible") is False:
            continue
        sub = agg[agg["sample"] == name].sort_values("time") if agg is not None and len(agg) else None
        if sub is None or sub.empty:
            continue

        color = st.get("color") or (pal[i] if i < len(pal) else "#333333")
        label = st.get("label") or name
        dash = st.get("dash") or DASHES[i % len(DASHES)] if cfg.get("vary_linetype") else (st.get("dash") or "solid")
        shape = st.get("shape") or (SHAPES[i % len(SHAPES)] if cfg.get("vary_shape", True) else "circle")

        t = sub["time"].to_numpy(dtype=float)
        y = sub["mean"].to_numpy(dtype=float)
        e = np.nan_to_num(sub["err"].to_numpy(dtype=float), nan=0.0) * mult

        mode = ("lines" if show_line else "") + ("+markers" if show_pts else "")
        mode = mode.strip("+") or "lines"

        if err_mode == "ribbon" and np.any(e > 0):
            lo, hi = y - e, y + e
            if cfg.get("y_scale_type") == "log":
                lo = np.maximum(lo, 1e-12)     # a log axis cannot show <=0
            fig.add_trace(go.Scatter(
                x=np.concatenate([t, t[::-1]]),
                y=np.concatenate([hi, lo[::-1]]),
                fill="toself", fillcolor=_rgba(color, float(cfg.get("shadow_alpha", 0.2))),
                line={"width": 0}, hoverinfo="skip", showlegend=False,
                name=f"{label} error", legendgroup=name))

        trace = go.Scatter(
            x=t, y=y, mode=mode, name=label, legendgroup=name,
            line={"color": color, "width": float(cfg.get("line_thickness", 2)), "dash": dash},
            marker={"color": color, "size": float(cfg.get("shape_size", 8)),
                    "symbol": shape,
                    "line": {"width": float(cfg.get("point_stroke", 1)), "color": color}},
            hovertemplate=(f"<b>{label}</b><br>%{{x}}<br>%{{y:.4g}}"
                           + ("±%{error_y.array:.3g}" if err_mode == "bar" else "")
                           + "<extra></extra>"),
        )
        if err_mode == "bar" and np.any(e > 0):
            trace.error_y = {
                "type": "data", "array": e, "visible": True,
                "thickness": float(cfg.get("error_bar_thickness", 1)),
                "width": float(cfg.get("error_bar_width", 4)),
                "color": color,
            }
        fig.add_trace(trace)

        if cfg.get("show_end_labels") and len(t):
            fig.add_annotation(x=t[-1], y=y[-1], text=label, showarrow=False,
                               xanchor="left", xshift=int(cfg.get("label_offset", 8)),
                               font={"size": cfg.get("label_font_size", 12), "color": color,
                                     "family": cfg.get("font_family", "Arial")})

    # ── highlight regions and time markers ──
    for r in (cfg.get("regions") or []):
        try:
            fig.add_vrect(x0=float(r["x0"]), x1=float(r["x1"]),
                          fillcolor=r.get("color", "#FFD54F"),
                          opacity=float(r.get("alpha", 0.2)), line_width=0,
                          annotation_text=r.get("label", ""), annotation_position="top left")
        except (KeyError, TypeError, ValueError):
            continue
    for m in (cfg.get("markers") or []):
        try:
            fig.add_vline(x=float(m["x"]), line_dash=m.get("dash", "dash"),
                          line_color=m.get("color", "#666666"),
                          line_width=float(m.get("width", 1)),
                          annotation_text=m.get("label", ""), annotation_position="top")
        except (KeyError, TypeError, ValueError):
            continue

    for ov in (analysis_overlays or []):
        fig.add_trace(ov)

    title = cfg.get("plot_title", "")
    subtitle = cfg.get("plot_subtitle", "")
    title_txt = title
    if subtitle:
        title_txt = (f"{title}<br><span style='font-size:{max(10, int(cfg.get('title_font_size', 18)) - 5)}px;"
                     f"color:#666'>{subtitle}</span>") if title else subtitle

    caption = ""
    if cfg.get("show_error_caption") and agg is not None and len(agg):
        et = {"sd": "SD", "sem": "SEM", "ci95": "95% CI"}.get(cfg.get("error_type", "sem"), "SEM")
        ns = sorted(set(int(v) for v in agg["n"].dropna().tolist())) if "n" in agg else []
        nstr = (f"n = {ns[0]}" if len(ns) == 1 else
                (f"n = {min(ns)}–{max(ns)}" if ns else ""))
        caption = f"Error bars: {et}" + (f"; {nstr}" if nstr else "")

    fig.update_layout(
        title={"text": title_txt,
               "font": {"size": cfg.get("title_font_size", 18),
                        "family": cfg.get("font_family", "Arial"),
                        "color": cfg.get("axis_color", "#000000")},
               "x": float(cfg.get("title_x", 0.5)), "xanchor": "center"},
        xaxis=_axis_kwargs(cfg, "x"),
        yaxis=_axis_kwargs(cfg, "y"),
        template=cfg.get("template", "simple_white"),
        font={"family": cfg.get("font_family", "Arial"),
              "size": cfg.get("axis_text_font_size", 14)},
        showlegend=bool(cfg.get("show_legend", True)),
        legend={"orientation": cfg.get("legend_orientation", "v"),
                "x": float(cfg.get("legend_x", 1.02)), "y": float(cfg.get("legend_y", 1)),
                "xanchor": "left", "yanchor": "top",
                "font": {"size": cfg.get("legend_font_size", 13),
                         "family": cfg.get("font_family", "Arial")},
                "bgcolor": "rgba(0,0,0,0)", "borderwidth": 0},
        margin={"l": 80, "r": 150 if cfg.get("show_legend", True) else 40, "t": 70, "b": 70},
        plot_bgcolor=cfg.get("plot_bgcolor", "#FFFFFF"),
        paper_bgcolor=cfg.get("paper_bgcolor", "#FFFFFF"),
        hovermode="closest",
    )
    if caption:
        fig.add_annotation(text=caption, xref="paper", yref="paper", x=0, y=-0.16,
                           showarrow=False, xanchor="left",
                           font={"size": max(9, int(cfg.get("axis_text_font_size", 14)) - 3),
                                 "color": "#555", "family": cfg.get("font_family", "Arial")})
    return fig


def _rgba(hexcolor, alpha):
    """#RRGGBB -> rgba() so ribbons can be translucent."""
    h = str(hexcolor).lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    try:
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    except (ValueError, IndexError):
        r, g, b = 100, 100, 100
    return f"rgba({r},{g},{b},{alpha})"


# ── Export ──────────────────────────────────────────────────────────────────
def figure_to_bytes(fig, fmt="png", width_in=10.0, height_in=8.0, dpi=300):
    """Render a figure to bytes at a real physical size.

    Kaleido works in pixels at scale=1, so inches x DPI is converted here. TIFF
    is produced by converting a high-DPI PNG, since kaleido has no TIFF writer.
    """
    fmt = (fmt or "png").lower()
    w_px = max(int(round(float(width_in) * 96)), 10)
    h_px = max(int(round(float(height_in) * 96)), 10)
    scale = max(float(dpi) / 96.0, 0.1)

    if fmt in ("html",):
        # Inline the whole plotly bundle rather than linking a CDN: the app's
        # promise is that it works offline and sends nothing to third parties,
        # and an exported figure emailed to a colleague must open the same way.
        return fig.to_html(include_plotlyjs=True, full_html=True).encode("utf-8")
    if fmt == "json":
        return fig.to_json().encode("utf-8")

    if fmt == "tiff":
        # kaleido has no TIFF writer, so tag a high-DPI PNG instead.
        png = render_images([fig], "png", width_in, height_in, dpi)[0]
        from PIL import Image
        im = Image.open(io.BytesIO(png))
        buf = io.BytesIO()
        # dpi must be tagged or the TIFF opens at the wrong physical size.
        im.save(buf, format="TIFF", dpi=(float(dpi), float(dpi)), compression="tiff_lzw")
        return buf.getvalue()

    return render_images([fig], fmt, width_in, height_in, dpi)[0]


def cumulative_frames(agg, cfg, styles, samples):
    """One figure per sample, each adding the next curve — the build animation.

    Axis ranges are frozen across frames so the plot doesn't jump around as
    curves appear, which is what made the R app's GIF hard to read.
    """
    frames = []
    if not samples:
        return frames
    fixed = dict(cfg)
    if agg is not None and len(agg):
        t = agg["time"].to_numpy(dtype=float)
        y = agg["mean"].to_numpy(dtype=float)
        e = np.nan_to_num(agg["err"].to_numpy(dtype=float), nan=0.0)
        if fixed.get("x_min") in (None, "") and fixed.get("x_max") in (None, "") and len(t):
            pad = (np.nanmax(t) - np.nanmin(t)) * 0.04 or 1
            fixed["x_min"], fixed["x_max"] = float(np.nanmin(t) - pad), float(np.nanmax(t) + pad)
        if fixed.get("y_min") in (None, "") and fixed.get("y_max") in (None, "") and len(y):
            lo, hi = np.nanmin(y - e), np.nanmax(y + e)
            if fixed.get("y_scale_type") == "log":
                lo = max(lo, np.nanmin(y[y > 0]) * 0.5 if np.any(y > 0) else 1e-3)
                fixed["y_min"], fixed["y_max"] = float(lo), float(hi * 1.6)
            else:
                pad = (hi - lo) * 0.08 or 0.1
                fixed["y_min"], fixed["y_max"] = float(lo - pad), float(hi + pad)
    for k in range(1, len(samples) + 1):
        frames.append(build_figure(agg, fixed, styles, samples=samples[:k]))
    return frames


# How long a static render may take before we give up and reclaim everything.
RENDER_TIMEOUT_S = 180


def _kill_tree(proc):
    """Kill a process AND its children.

    The render worker spawns a headless Chrome. Terminating only the Python
    child would orphan the browser — the exact leak this indirection exists to
    prevent — so on Windows use taskkill /T, elsewhere kill the process group.
    """
    import os
    import signal
    import subprocess
    try:
        if os.name == "nt":
            subprocess.run(["taskkill", "/F", "/T", "/PID", str(proc.pid)],
                           capture_output=True, timeout=30)
        else:
            os.killpg(os.getpgid(proc.pid), signal.SIGKILL)
    except Exception:
        try:
            proc.kill()
        except Exception:
            pass


def render_images(figs, fmt, width_in, height_in, dpi, timeout=RENDER_TIMEOUT_S):
    """Render figures to bytes in a throwaway subprocess.

    Out-of-process for three reasons: a hung browser cannot wedge the app, the
    OS reclaims every child on exit (in-process rendering leaked one Chrome per
    call), and one subprocess renders the whole batch through plotly's batch
    writer rather than paying ~12 s of browser start-up per frame.
    """
    import json
    import os
    import subprocess
    import sys
    import tempfile

    if not figs:
        return []
    w_px = max(int(round(float(width_in) * 96)), 10)
    h_px = max(int(round(float(height_in) * 96)), 10)
    # Vector formats are resolution-independent; scaling them only changes the
    # nominal page size, so pin scale to 1 for them.
    scale = 1.0 if fmt in ("pdf", "svg") else max(float(dpi) / 96.0, 0.1)

    with tempfile.TemporaryDirectory() as td:
        paths = [os.path.join(td, f"f{i:04d}.{fmt}") for i in range(len(figs))]
        spec = os.path.join(td, "spec.json")
        with open(spec, "w", encoding="utf-8") as fh:
            json.dump({"figs": [f.to_plotly_json() if hasattr(f, "to_plotly_json") else f
                                for f in figs],
                       "paths": paths, "format": fmt,
                       "width": w_px, "height": h_px, "scale": scale}, fh)

        root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        # NB: capture_output is a subprocess.run() kwarg — Popen needs the pipes
        # named explicitly.
        kw = {"cwd": root, "stdout": subprocess.PIPE, "stderr": subprocess.PIPE,
              "text": True}
        if os.name != "nt":
            kw["start_new_session"] = True          # own process group to kill
        proc = subprocess.Popen([sys.executable, "-m", "app.render_worker", spec], **kw)
        try:
            _out, err = proc.communicate(timeout=timeout)
        except subprocess.TimeoutExpired:
            _kill_tree(proc)
            raise TimeoutError(
                f"Rendering took longer than {timeout}s and was cancelled. "
                "Try a smaller size or DPI, or use the PNG/SVG buttons, which "
                "render in your browser and need no external engine.")
        if proc.returncode != 0:
            msg = (err or "").strip().splitlines()[-1] if err else "unknown error"
            # The engine drives a headless Chrome. The usual causes are no Chrome
            # installed, or a Chrome that just auto-updated. Say what to do
            # instead of surfacing a library traceback.
            if "browser" in msg.lower() or "chrome" in msg.lower():
                raise RuntimeError(
                    "This format needs a headless browser and one could not be "
                    "started. The PNG, JPEG and SVG buttons work without it "
                    "(they render in this browser). To enable PDF/TIFF/"
                    "PowerPoint/GIF, install a private copy for the app by "
                    r"running:  .venv\Scripts\plotly_get_chrome")
            raise RuntimeError(f"Image export failed: {msg}")

        out = []
        for p in paths:
            with open(p, "rb") as fh:
                out.append(fh.read())
        return out


def render_many(figs, width_in, height_in, dpi):
    """PNG bytes for a list of figures (used by the PPTX and GIF builders)."""
    return render_images(figs, "png", width_in, height_in, dpi)


def export_pptx(agg, cfg, styles, samples, width_in=10.0, height_in=7.5, dpi=200):
    """PowerPoint deck: one slide per cumulative frame, image-per-slide."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(float(width_in))
    prs.slide_height = Inches(float(height_in))
    blank = prs.slide_layouts[6]
    figs = cumulative_frames(agg, cfg, styles, samples)
    for png in render_many(figs, width_in, height_in, dpi):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(io.BytesIO(png), 0, 0,
                                 width=prs.slide_width, height=prs.slide_height)
    if not len(prs.slides._sldIdLst):
        prs.slides.add_slide(blank)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def export_gif(agg, cfg, styles, samples, width_in=10.0, height_in=8.0,
               dpi=100, fps=1.0, hold_last=3):
    """Animated GIF of the cumulative build, holding the final frame."""
    from PIL import Image

    figs = cumulative_frames(agg, cfg, styles, samples)
    if not figs:
        raise ValueError("nothing to animate")
    imgs = [Image.open(io.BytesIO(b)).convert("P", palette=Image.ADAPTIVE)
            for b in render_many(figs, width_in, height_in, dpi)]
    imgs.extend([imgs[-1]] * max(0, int(hold_last)))   # linger on the full plot
    buf = io.BytesIO()
    imgs[0].save(buf, format="GIF", save_all=True, append_images=imgs[1:],
                 duration=int(1000 / max(float(fps), 0.1)), loop=0, optimize=True)
    return buf.getvalue()
